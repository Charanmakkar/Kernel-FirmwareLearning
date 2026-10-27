from pathlib import Path
import csv
import hashlib
import re
from urllib.parse import unquote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]/'Embedded_Linux_Kernel_Firmware_Learning'
OUT=ROOT/'Word_Edition'
OUT.mkdir(exist_ok=True)
BLUE='16324F';TEAL='007F86';GREY='526474'

def bookmark_name(path):
    return 'f_'+hashlib.sha1(str(path).encode()).hexdigest()[:20]

def field(paragraph,instruction):
    run=paragraph.add_run()
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    code=OxmlElement('w:instrText');code.set(qn('xml:space'),'preserve');code.text=instruction
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
    text=OxmlElement('w:t');text.text=('1' if instruction.strip()=='PAGE' else 'Refresh this table in Word (Ctrl+A, F9). Headings are available in the Navigation Pane.')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for node in [begin,code,separate,text,end]:run._r.append(node)

def shade(cell,color):
    props=cell._tc.get_or_add_tcPr()
    element=OxmlElement('w:shd');element.set(qn('w:fill'),color);props.append(element)

def link(paragraph,label,target,source,included):
    hyperlink=OxmlElement('w:hyperlink')
    if target.startswith(('https://','http://')):
        rid=paragraph.part.relate_to(target,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        hyperlink.set(qn('r:id'),rid)
    else:
        resolved=(source.parent/unquote(target.split('#')[0])).resolve() if target.split('#')[0] else source
        if resolved in included:
            hyperlink.set(qn('w:anchor'),bookmark_name(resolved))
        else:
            try:rel='../'+resolved.relative_to(ROOT).as_posix()
            except ValueError:rel=target
            rid=paragraph.part.relate_to(rel,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            hyperlink.set(qn('r:id'),rid)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),TEAL);props.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
    run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text)
    hyperlink.append(run);paragraph._p.append(hyperlink)

pattern=re.compile(r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|`[^`]+`)')
def inline(paragraph,text,source,included):
    for token in pattern.split(text):
        if not token:continue
        match=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',token)
        if match:link(paragraph,match[1],match[2],source,included)
        elif token.startswith('**') and token.endswith('**'):paragraph.add_run(token[2:-2]).bold=True
        elif token.startswith('`') and token.endswith('`'):
            r=paragraph.add_run(token[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        else:paragraph.add_run(token)

def add_bookmark(paragraph,path,n):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(n));start.set(qn('w:name'),bookmark_name(path))
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(n))
    paragraph._p.insert(0,start);paragraph._p.append(end)

def render_markdown(doc,path,included,n):
    lines=path.read_text(encoding='utf-8').splitlines()
    i=0;first=True
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].startswith('```'):
                p=doc.add_paragraph(style='Code');p.add_run(lines[i].expandtabs(4) or ' ')
                i+=1
            i+=1;continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[ :\-|]+\|$',lines[i+1]):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[ :\-|]+\|$',lines[i]):rows.append([x.strip() for x in lines[i].strip('|').split('|')])
                i+=1
            if not rows:continue
            count=max(len(r) for r in rows)
            t=doc.add_table(rows=1,cols=count);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
            for j,celltext in enumerate(rows[0]):
                inline(t.rows[0].cells[j].paragraphs[0],celltext,path,included)
                shade(t.rows[0].cells[j],BLUE)
                for run in t.rows[0].cells[j].paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row_n,data in enumerate(rows[1:]):
                cells=t.add_row().cells
                for j,celltext in enumerate(data):inline(cells[j].paragraphs[0],celltext,path,included)
                if row_n%2==0:
                    for cell in cells:shade(cell,'F1F5F8')
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(4)
                        for run in p.runs:run.font.size=Pt(8.5)
            doc.add_paragraph();continue
        heading=re.match(r'^(#{1,6})\s+(.*)',line)
        if heading:
            level=min(len(heading[1]),4)
            p=doc.add_paragraph(style=f'Heading {level}');inline(p,heading[2],path,included)
            if first:add_bookmark(p,path,n);first=False
            i+=1;continue
        bullet=re.match(r'^([-*]|\d+\.)\s+(.*)',line)
        if bullet:
            p=doc.add_paragraph(style='List Bullet' if bullet[1] in '-*' else 'List Number')
            inline(p,bullet[2],path,included);i+=1;continue
        if line.startswith('> '):
            p=doc.add_paragraph(style='Quote');inline(p,line[2:],path,included);i+=1;continue
        paragraph=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|```|\||[-*] |\d+\. |> )',lines[i]):
            paragraph.append(lines[i]);i+=1
        p=doc.add_paragraph();inline(p,' '.join(paragraph),path,included)

def make(title,subtitle,paths,filename,extras=None):
    doc=Document();sec=doc.sections[0]
    sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
    sec.top_margin=sec.bottom_margin=Inches(.72)
    sec.left_margin=sec.right_margin=Inches(.72)
    styles=doc.styles
    normal=styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.1
    for name,size,color in [('Title',32,BLUE),('Subtitle',14,GREY),('Heading 1',21,BLUE),('Heading 2',15,TEAL),('Heading 3',12,BLUE),('Heading 4',11,BLUE)]:
        styles[name].font.name='Calibri';styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor.from_string(color)
        styles[name].paragraph_format.keep_with_next=True
    styles['Heading 1'].paragraph_format.page_break_before=True
    if 'Code' not in styles:styles.add_style('Code',1)
    code=styles['Code'];code.font.name='Consolas';code.font.size=Pt(8)
    code.paragraph_format.space_after=Pt(0);code.paragraph_format.line_spacing=1.0
    code.paragraph_format.left_indent=Inches(.12)
    for name in ['Normal','Code','List Bullet','List Number']:styles[name].paragraph_format.widow_control=True
    # Keep a small header, explicit staged status, and a real Word page-number field.
    h=sec.header.paragraphs[0];h.text='EMBEDDED LINUX  /  KERNEL  /  FIRMWARE'
    h.runs[0].font.size=Pt(8);h.runs[0].font.color.rgb=RGBColor.from_string(GREY)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Release 1 • full program incomplete  |  ').font.size=Pt(8)
    field(footer,' PAGE ')
    doc.add_paragraph('PRACTICAL LEARNING & INTERVIEW PREPARATION',style='Subtitle')
    doc.add_paragraph(title,style='Title')
    doc.add_paragraph(subtitle,style='Subtitle')
    doc.add_paragraph('27 October 2026  •  Junior-to-mid-level preparation  •  10–12 study hours/week')
    box=doc.add_table(rows=1,cols=1);shade(box.cell(0,0),'E8F3F3')
    box.cell(0,0).text=('STAGED RELEASE: The full master prompt is not complete. This release contains one complete lesson and one tested project, all 15 project plans, the roadmap and a 200-question bank. Remaining lessons, 14 implementations, later assessment fixtures and physical validation are pending.')
    doc.add_paragraph('Use Markdown for edits, CSVs for tracking, source files for builds, and evidence JSON for executed results. The supplied tests do not assess your skill.')
    doc.add_paragraph('Contents and navigation',style='Heading 1')
    doc.add_paragraph('Use Word’s Navigation Pane to browse headings. The table below updates when fields are refreshed in Word. Cross-references within this volume open the corresponding document section; links to other package files resolve relative to Word_Edition/.')
    field(doc.add_paragraph(),' TOC \\o "1-2" \\h \\z \\u ')
    included={p.resolve() for p in paths}
    for n,path in enumerate(paths,1):render_markdown(doc,path.resolve(),included,n)
    if extras:extras(doc)
    doc.core_properties.title=title
    doc.core_properties.subject='Release 1: practical embedded learning program; incomplete full master program'
    doc.core_properties.author='Codex — educational reference material'
    doc.core_properties.keywords='Embedded Linux; C; kernel; drivers; firmware; RTOS; release 1'
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    doc.save(OUT/filename)
    print(filename, 'paragraphs=',len(doc.paragraphs),'tables=',len(doc.tables))

main=[
'00_Start_Here/PROFILE_AND_ASSUMPTIONS.md','00_Start_Here/ENVIRONMENT.md',
'01_Learning_Roadmap/PREREQUISITES_AND_GATES.md','03_JD_JR_Skills_and_Learning_Content/JD_ANALYSIS_AND_GAP_CLOSURE.md',
'03_JD_JR_Skills_and_Learning_Content/MODULE_CATALOG.md','03_JD_JR_Skills_and_Learning_Content/M01_C_Objects_Bytes_and_Contracts.md',
'01_Learning_Roadmap/PROJECT_CATALOG.md']
with (ROOT/'01_Learning_Roadmap/project_plan.csv').open(encoding='utf-8-sig') as f:
    main += [r['Folder']+'/README.md' for r in csv.DictReader(f)]
main += ['07_Projects/01_Basic/B01_Embedded_C_Utility_Library/VERSIONS.md',
'07_Projects/01_Basic/B01_Embedded_C_Utility_Library/docs/REVIEW.md',
'07_Projects/04_Capstone/CAPSTONE_REQUIREMENTS.md','04_Skill_Assessments/ASSESSMENT_PLAN.md',
'05_Interview_Strategy/INTERVIEW_PLAYBOOK.md','06_AI_Integration/AI_ASSISTED_ENGINEERING.md',
'06_AI_Integration/OPTIONAL_EMBEDDED_ANOMALY_EXTENSION.md','08_Portfolio_and_Job_Readiness/PORTFOLIO_EVIDENCE.md',
'09_References_and_Resources/REFERENCES.md','10_Progress_Tracking/VALIDATION.md','PROGRESS.md']

def appendices(doc):
    for csvpath,title,fields in [
      ('01_Learning_Roadmap/roadmap_64_weeks.csv','Week-by-week roadmap',['Week','Phase','Learning objectives','Practical exercise','Objective exit test']),
      ('03_JD_JR_Skills_and_Learning_Content/reference_skills_matrix.csv','Reference skills summary',['Requirement ID','Role track','Requirement','Expected depth','Learning module']),
      ('05_Interview_Strategy/revision_30_days.csv','30-day revision calendar',['Day','Focus','Independent task','Links/IDs'])]:
        doc.add_paragraph(title,style='Heading 1')
        doc.add_paragraph('Editable full detail is supplied in '+csvpath+'. All skills begin not assessed; future module links identify planned material.')
        with (ROOT/csvpath).open(encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
        for row in rows:
            doc.add_paragraph(' • '.join(row[x] for x in fields[:2]),style='Heading 3')
            for key in fields[2:]:
                p=doc.add_paragraph();p.add_run(key+': ').bold=True;p.add_run(row[key])
    doc.add_paragraph('B01 complete source appendix',style='Heading 1')
    doc.add_paragraph('The source files in the package are authoritative for builds. This appendix makes the first project self-contained for Word reading. Do not retype wrapped code from a document for validation.')
    project=ROOT/'07_Projects/01_Basic/B01_Embedded_C_Utility_Library'
    for rel in ['include/embedded_util.h','src/embedded_util.c','src/demo.c','tests/test_util.c','scripts/verify.py','Makefile']:
        doc.add_paragraph(rel,style='Heading 2')
        for line in (project/rel).read_text(encoding='utf-8').splitlines():doc.add_paragraph(line.expandtabs(4) or ' ',style='Code')

make('Embedded Linux, Kernel & Firmware','Learning program • first complete lesson • 15-project plan • tested C reference',
     [ROOT/p for p in main],'Learning_Program_Release_1.docx',appendices)
bank=[ROOT/'02_Interview_Questions_and_Answers/README.md']+sorted((ROOT/'02_Interview_Questions_and_Answers').glob('[0-9][0-9][0-9]_*.md'))
make('200 Interview Questions & Answers','Three tracks • spoken and deeper answers • coding, debugging and design',bank,'Interview_Bank_200_Release_1.docx')
questions=[ROOT/p for p in ['04_Skill_Assessments/questions/D01_Initial_Diagnostic.md',
 '03_JD_JR_Skills_and_Learning_Content/exercises/M01_Questions.md',
 '07_Projects/01_Basic/B01_Embedded_C_Utility_Library/exercises/EXERCISES.md',
 '02_Interview_Questions_and_Answers/questions_only.md']]
make('Assessment & Practice Questions','Independent attempts • keep answer documents closed while working',questions,'Assessment_Questions_Release_1.docx')
keys=[ROOT/p for p in ['04_Skill_Assessments/answer_keys/D01_Key.md',
 '03_JD_JR_Skills_and_Learning_Content/solutions/M01_Solutions.md',
 '07_Projects/01_Basic/B01_Embedded_C_Utility_Library/exercises/REFERENCE_SOLUTIONS.md']]
make('Assessment Answer Keys','D01 diagnostic • M01 exit exercise • B01 independent extensions',keys,'Assessment_Answer_Keys_Release_1.docx')
